"""Round-trip check: xlsx -> pdf -> xlsx keeps the table. Run: python test_convert.py"""
import tempfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

from convert import excel_to_pdf, pdf_to_excel

with tempfile.TemporaryDirectory() as d:
    src, pdf, back = (str(Path(d) / n) for n in ("in.xlsx", "mid.pdf", "out.xlsx"))
    wb = Workbook()
    rows = [["Name", "Qty", "Price"]] + [[f"Item <{i}> & co", i, i * 1.5] for i in range(1, 80)]
    for r in rows:
        wb.active.append(r)
    wb.save(src)

    excel_to_pdf(src, pdf)
    pdf_to_excel(pdf, back)

    out = [list(r) for ws in load_workbook(back).worksheets for r in ws.iter_rows(values_only=True)]
    data = [r for r in out if r[0] != "Name"]  # header repeats on each PDF page
    assert out[0] == ["Name", "Qty", "Price"], out[0]
    assert data == rows[1:], (data[:3], rows[1:4])
    print(f"OK: {len(data)} rows round-tripped")

# statement parsing: wrapped tail line, description split above its date line, balance check
from reportlab.pdfgen import canvas

from convert import parse_statement, pdf_to_csv

lines = ["Account Number: 000123", "January 01, 2025 through March 31, 2025",
         "Beginning Balance $100.00", "Fees -4.95", "Ending Balance $1,095.05",
         "DATE DESCRIPTION AMOUNT BALANCE", "Beginning Balance $100.00",
         "01/02 Card Purchase 01/02 Shop IL Card -10.00 90.00", "1425",
         "01/03Online", "01/03 Transfer From Chk 1,010.00 1,100.00",
         "01/31 Monthly Service Fee -4.95 1,095.05", "Ending Balance $1,095.05", "Page 1 of 1"]
with tempfile.TemporaryDirectory() as d:
    pdf, out = str(Path(d) / "st.pdf"), str(Path(d) / "st.csv")
    c = canvas.Canvas(pdf)
    for i, line in enumerate(lines):
        c.drawString(40, 800 - 16 * i, line)
    c.save()
    st = parse_statement(pdf)
    assert st["transactions"] == [["01/02", "Card Purchase 01/02 Shop IL Card 1425", -10.0, 90.0],
                                  ["01/03", "Online Transfer From Chk", 1010.0, 1100.0],
                                  ["01/31", "Monthly Service Fee", -4.95, 1095.05]], st["transactions"]
    assert st["info"] == {"Account Number": "000123", "Statement Period": "January 01, 2025 - March 31, 2025"}
    assert st["summary"]["Fees"] == -4.95 and st["mismatches"] == [], st
    pdf_to_csv(pdf, out)
    assert open(out, encoding="utf-8-sig").readline().strip() == "Date,Description,Amount,Balance"
    print("OK: statement parsed and reconciled")

# Missing PDF Unicode mappings must be recovered from the correct font, or
# retained visibly. Silently dropping a glyph can alter a date or account ID.
from types import SimpleNamespace
from unittest.mock import patch
from convert import _page_text

class GlyphPage:
    def __init__(self):
        self.chars = [{"fontname": "Demo+Calibri", "text": value}
                      for value in ["0", "(cid:1007)", "/28 Chk", "(cid:856)", "5961"]]
        self.page_obj = SimpleNamespace(resources={"Font": {
            "F1": {"BaseFont": "Demo+Calibri"}}})

    def extract_text(self):
        return "".join(char["text"] for char in self.chars)

with patch("convert._embedded_glyph_map", return_value={1007: "3", 856: "."}):
    assert _page_text(GlyphPage()) == "03/28 Chk.5961"
with patch("convert._embedded_glyph_map", return_value={}):
    assert "(cid:1007)" in _page_text(GlyphPage())
print("OK: unmapped glyphs recovered or preserved, never silently deleted")
