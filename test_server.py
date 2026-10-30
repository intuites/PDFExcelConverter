"""End-to-end check of server.py: xlsx -> pdf -> xlsx over HTTP, plus bad input. Run: python test_server.py"""
import io
import threading
import urllib.error
import urllib.request
from functools import partial
from http.server import ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook

from server import Handler

srv = ThreadingHTTPServer(("127.0.0.1", 0), partial(Handler, directory=str(Path(__file__).parent / "web")))
threading.Thread(target=srv.serve_forever, daemon=True).start()
base = f"http://127.0.0.1:{srv.server_port}"


def post(name, body):
    return urllib.request.urlopen(urllib.request.Request(f"{base}/convert?name={name}", data=body, method="POST"))


wb = Workbook()
for r in [["A", "B"], ["x", 1], ["y", 2]]:
    wb.active.append(r)
buf = io.BytesIO()
wb.save(buf)

pdf = post("data.xlsx", buf.getvalue())
assert pdf.headers["Content-Disposition"] == 'attachment; filename="data.pdf"'
xlsx = post("data.pdf", pdf.read())
rows = [list(r) for r in load_workbook(io.BytesIO(xlsx.read())).active.iter_rows(values_only=True)]
assert rows == [["A", "B"], ["x", 1], ["y", 2]], rows

for name, body, code in [("a.txt", b"hi", 400), ("a.pdf", b"not a pdf", 422)]:
    try:
        post(name, body)
        raise AssertionError(f"{name} should fail")
    except urllib.error.HTTPError as e:
        assert e.code == code, (name, e.code)

assert b"PDF" in urllib.request.urlopen(base + "/").read()
srv.shutdown()
print("OK: server round-trip and error cases pass")
